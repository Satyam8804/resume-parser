import json
import logging
import os
from datetime import datetime
from io import BytesIO
from pathlib import Path

from dateutil import parser as date_parser
from dateutil.relativedelta import relativedelta
from docx import Document
from dotenv import load_dotenv
from fastapi import APIRouter, File, HTTPException, UploadFile
from groq import Groq
from pydantic import ValidationError
from pypdf import PdfReader

from app.schemas.user import User

load_dotenv()

logger = logging.getLogger("uvicorn.error")

router = APIRouter()

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

groq_api_key = os.getenv("GROQ_API_KEY")

if not groq_api_key:
    raise ValueError("GROQ_API_KEY is not set")

# Explicit timeout so a slow/unreachable Groq call fails fast instead of hanging.
client = Groq(api_key=groq_api_key, timeout=45.0, max_retries=1)

model = "llama-3.3-70b-versatile"


def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    # Join pages with a newline so words at page boundaries don't get glued together.
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text_from_docx(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts = []

    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # also pull text out of tables, since resumes sometimes use them for layout
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def extract_text(data: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        return extract_text_from_pdf(data)
    elif suffix == ".docx":
        return extract_text_from_docx(data)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {suffix}")


CURRENT_JOB_MARKERS = {"present", "current", "currently working", "ongoing", "till date", "now", ""}


def compute_years(start_date: str, end_date: str) -> float | None:
    if not start_date:
        return None
    try:
        start = date_parser.parse(start_date, default=datetime(1900, 1, 1), fuzzy=True)
    except (ValueError, TypeError, OverflowError):
        return None

    end_normalized = (end_date or "").strip().lower()
    if end_normalized in CURRENT_JOB_MARKERS:
        end = datetime.now()
    else:
        try:
            end = date_parser.parse(end_date, default=datetime(1900, 1, 1), fuzzy=True)
        except (ValueError, TypeError, OverflowError):
            end = datetime.now()

    if end < start:
        return None

    delta = relativedelta(end, start)
    total_months = delta.years * 12 + delta.months
    if total_months < 0:
        return None

    return round(total_months / 12, 1)


# Plain `def` (not `async def`): FastAPI runs it in a threadpool, so the blocking
# PDF parsing and the synchronous Groq call don't freeze the event loop.
@router.post("/api/upload")
def upload_file(file: UploadFile = File(...)):
    file_suffix = Path(file.filename or "").suffix.lower()

    if file_suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{file_suffix}'. Only PDF and DOCX are supported.",
        )

    data = file.file.read()

    try:
        text = extract_text(data, file_suffix)
    except HTTPException:
        raise
    except Exception:
        logger.exception("Text extraction failed for %s", file.filename)
        raise HTTPException(status_code=422, detail="Could not read this file. It may be corrupted or password-protected.")

    if not text or not text.strip():
        raise HTTPException(status_code=422, detail="Could not extract any text from the uploaded file.")

    schema = User.model_json_schema()

    response_format = {"type": "json_object"}

    system_prompt = f"""
    You are an expert resume information extraction assistant.
    Your task is to extract structured information from the resume provided by the user.
    Rules:
    1. Return ONLY valid JSON.
    2. The JSON must strictly follow the provided schema.
    3. Do not include any explanations, markdown, or extra text.
    4. If a field is missing or cannot be determined, return null for that field.
    5. Do not invent or infer information that is not explicitly mentioned in the resume.
    6. Preserve the original spelling and capitalization of names, companies, skills, and job titles.
    7. Extract all work experiences in chronological order if possible.
    8. Extract all technical and non-technical skills mentioned.
    9. If the resume contains multiple experiences or skills, include all of them in the corresponding arrays.
    10. For each experience, extract start_date exactly as written (e.g. "Jan 2023", "2021").
        If the candidate is currently working at that job (indicated by words like
        "Present", "Current", "Ongoing", "Till date", "Now", or an empty/missing end date
        for the most recent role), set end_date to the literal string "Present".
        Otherwise extract end_date exactly as written.
    11. Do NOT calculate or guess years_of_experience yourself — leave it null; it will be
        computed separately from the dates.

    Return JSON strictly matching the following schema:

    {schema}
    """

    message_system = {"role": "system", "content": system_prompt}

    prompts = f'''This is a candidate resume. please extract all information and from this {text}'''

    message = {"role": "user", "content": prompts}

    messages = [message_system, message]

    # NOTE: LLM failures return 500, not 502. A 502 in the browser should only ever
    # mean Render's proxy couldn't reach the app, never "my own code said so".
    try:
        response = client.chat.completions.create(
            model=model,
            messages=messages,
            response_format=response_format,
        )
    except Exception as e:
        logger.exception("Groq request failed")
        raise HTTPException(status_code=500, detail=f"LLM request failed: {e}")

    answer = response.choices[0].message.content

    try:
        data_file = json.loads(answer)
    except (json.JSONDecodeError, TypeError):
        logger.error("LLM returned malformed JSON: %.500s", answer)
        raise HTTPException(status_code=500, detail="LLM returned malformed JSON. Please retry.")

    try:
        user = User(**data_file)
    except ValidationError as e:
        raise HTTPException(status_code=422, detail=f"Extracted data did not match expected schema: {e.errors()}")

    for exp in user.experiences:
        exp.years_of_experience = compute_years(exp.start_date, exp.end_date)

    valid_durations = [exp.years_of_experience for exp in user.experiences if exp.years_of_experience is not None]
    total_experience = round(sum(valid_durations), 1) if valid_durations else None

    user.total_experience = total_experience

    return {"data": user}